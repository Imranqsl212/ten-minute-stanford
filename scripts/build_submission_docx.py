from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "submission" / "article_draft.md"
OUTPUT = ROOT / "submission" / "article_draft.docx"
LINK_RE = re.compile(r"\[([^]]+)]\((https?://[^)]+)\)")
IMAGE_RE = re.compile(r"^!\[([^]]*)]\(([^)]+)\)$")
CASE_HEADINGS = (
    "Near the median",
    "Near the limit",
    "Above the limit",
    "Longest modeled example",
)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "8C1515")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_rich_text(paragraph, text: str) -> None:
    cursor = 0
    for match in LINK_RE.finditer(text):
        paragraph.add_run(text[cursor : match.start()])
        add_hyperlink(paragraph, match.group(1), match.group(2))
        cursor = match.end()
    paragraph.add_run(text[cursor:])


def shade_paragraph(paragraph, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def keep_table_row_together(row) -> None:
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def add_case_grid(document: Document, cards: list[tuple[str, list[str]]]) -> None:
    table = document.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"

    for row in table.rows:
        keep_table_row_together(row)

    for index, (title, lines) in enumerate(cards):
        cell = table.cell(index // 2, index % 2)
        cell.width = Inches(3.1)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        shade_cell(cell, "F7F4F1")
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0

        title_run = paragraph.add_run(title)
        title_run.bold = True
        title_run.font.name = "Arial"
        title_run.font.size = Pt(10.5)
        title_run.font.color.rgb = RGBColor(140, 21, 21)

        for line_index, line in enumerate(lines):
            run = paragraph.add_run(f"\n{line}")
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            run.bold = line_index == 0

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def parse_source() -> tuple[dict[str, str], list[str]]:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    metadata: dict[str, str] = {}
    body: list[str] = []
    in_content = False

    for line in lines:
        stripped = line.strip()
        if stripped == "CONTENT":
            in_content = True
            continue
        if not in_content:
            if not stripped:
                continue
            key, value = stripped.split(":", maxsplit=1)
            metadata[key] = value.strip()
        else:
            body.append(line)

    return metadata, body


def build_document() -> Path:
    metadata, body = parse_source()
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = document.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(10.75)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color in (
        ("Heading 1", 17, RGBColor(33, 33, 33)),
        ("Heading 2", 13, RGBColor(140, 21, 21)),
    ):
        style = document.styles[style_name]
        style.font.name = "Georgia"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.keep_with_next = True

    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run(metadata["TITLE"])
    title_run.font.name = "Georgia"
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(33, 33, 33)

    table = document.add_table(rows=4, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.0)
    table.columns[1].width = Inches(5.8)
    for row, key in zip(table.rows, ("SLUG", "AUTHORS", "SECTION", "STATUS"), strict=True):
        row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        label = row.cells[0].paragraphs[0].add_run(key)
        label.bold = True
        label.font.name = "Arial"
        label.font.size = Pt(8)
        label.font.color.rgb = RGBColor(140, 21, 21)
        value = row.cells[1].paragraphs[0].add_run(metadata[key])
        value.font.name = "Arial"
        value.font.size = Pt(9)

    document.add_paragraph()

    paragraphs: list[str] = []
    buffer: list[str] = []
    for line in body:
        if line.strip():
            buffer.append(line.strip())
        elif buffer:
            paragraphs.append(" ".join(buffer))
            buffer = []
    if buffer:
        paragraphs.append(" ".join(buffer))

    index = 0
    while index < len(paragraphs):
        text = paragraphs[index]
        if text == f"### {CASE_HEADINGS[0]}":
            cards: list[tuple[str, list[str]]] = []
            for heading_index, heading in enumerate(CASE_HEADINGS):
                expected = f"### {heading}"
                if paragraphs[index] != expected:
                    raise ValueError(f"Expected course-card heading {expected!r}")
                next_heading = (
                    f"### {CASE_HEADINGS[heading_index + 1]}"
                    if heading_index + 1 < len(CASE_HEADINGS)
                    else None
                )
                card_lines: list[str] = []
                index += 1
                while index < len(paragraphs):
                    candidate = paragraphs[index]
                    if candidate == next_heading:
                        break
                    if next_heading is None and candidate.startswith(
                        "Examples were selected by a fixed rule"
                    ):
                        break
                    card_lines.append(candidate)
                    index += 1
                cards.append((heading, card_lines))
            add_case_grid(document, cards)
            continue

        if text.startswith(("## ", "### ")):
            level = 2 if text.startswith("### ") else 1
            heading_text = text[4:] if level == 2 else text[3:]
            document.add_heading(heading_text, level=level)
            index += 1
            continue

        image_match = IMAGE_RE.fullmatch(text)
        if image_match:
            image_path = (SOURCE.parent / image_match.group(2)).resolve()
            if not image_path.is_file():
                raise FileNotFoundError(f"Missing submission figure: {image_path}")
            document.add_picture(str(image_path), width=Inches(6.3))
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            document.paragraphs[-1].paragraph_format.keep_with_next = True
            index += 1
            continue

        paragraph = document.add_paragraph()
        if text.startswith("CAPTION: "):
            add_rich_text(paragraph, text.removeprefix("CAPTION: "))
            for run in paragraph.runs:
                run.italic = True
                run.font.size = Pt(9)
            paragraph.paragraph_format.space_after = Pt(10)
            index += 1
            continue
        if text.startswith("CALLOUT: "):
            shade_paragraph(paragraph, "F2EDE7")
            add_rich_text(paragraph, text.removeprefix("CALLOUT: "))
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(9.5)
                run.bold = True
            index += 1
            continue
        if text.startswith("LABEL: "):
            run = paragraph.add_run(text.removeprefix("LABEL: ").upper())
            run.font.name = "Arial"
            run.font.size = Pt(8)
            run.font.bold = True
            run.font.color.rgb = RGBColor(140, 21, 21)
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_after = Pt(2)
            index += 1
            continue
        add_rich_text(paragraph, text)
        if index == 0:
            paragraph.paragraph_format.space_after = Pt(14)
            for run in paragraph.runs:
                run.font.name = "Georgia"
                run.font.size = Pt(13)
                run.font.color.rgb = RGBColor(63, 59, 56)
        index += 1

    document.core_properties.title = metadata["TITLE"]
    document.core_properties.author = metadata["AUTHORS"]
    document.core_properties.subject = "The Stanford Daily Data submission draft"
    document.core_properties.keywords = "Stanford, course scheduling, walking, data"
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
